import builtins
import json
import zipfile
from pathlib import Path

from scanner.parsers import (STATUS_ENCRYPTED, STATUS_METADATA_ONLY, STATUS_OK, STATUS_PARTIAL, STATUS_UNREADABLE,
                             STATUS_UNSUPPORTED, ContentExtractor)


def _minimal_pdf(text: str) -> bytes:
    """テキストを 1 行持つ最小限の PDF を生成する。"""
    stream = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def test_text_markdown_json_yaml_xml(tmp_path):
    ex = ContentExtractor()
    for name, body in [("a.txt", "hello 世界"), ("b.md", "# 見出し"), ("c.json", json.dumps({"k": "値"})),
                       ("d.yaml", "key: value"), ("e.xml", "<root><a>x</a></root>"), ("f.log", "ERROR boot")]:
        p = tmp_path / name
        p.write_text(body, encoding="utf-8")
        r = ex.extract(p)
        assert r.status == STATUS_OK, name
        assert body.split()[0].strip("#{") in r.text or "値" in r.text


def test_cp932_text(tmp_path):
    p = tmp_path / "sjis.txt"
    p.write_bytes("議事録 会議".encode("cp932"))
    assert "議事録" in ContentExtractor().extract(p).text


def test_csv_and_html(tmp_path):
    ex = ContentExtractor()
    c = tmp_path / "t.csv"
    c.write_text("name,score\nalpha,10\n", encoding="utf-8")
    r = ex.extract(c)
    assert r.status == STATUS_OK and "alpha" in r.text and r.metadata["header"] == "name score"
    h = tmp_path / "t.html"
    h.write_text("<html><title>報告</title><script>evil()</script><body><p>本文</p></body></html>", encoding="utf-8")
    r = ex.extract(h)
    assert "本文" in r.text and "evil" not in r.text and r.metadata["title"] == "報告"


def test_source_code_is_read_not_executed(tmp_path):
    marker = tmp_path / "executed.flag"
    code = tmp_path / "danger.py"
    code.write_text(f"open(r'{marker}', 'w').write('x')\nprint('motor control')\n", encoding="utf-8")
    r = ContentExtractor().extract(code)
    assert r.kind == "code" and "motor control" in r.text
    assert not marker.exists()


def test_large_file_is_truncated(tmp_path):
    p = tmp_path / "big.log"
    p.write_text("line\n" * 100000, encoding="utf-8")
    r = ContentExtractor(max_bytes=1000, max_chars=500).extract(p)
    assert r.status == STATUS_PARTIAL and len(r.text) <= 500


def test_pdf_text_and_encrypted(tmp_path):
    from pypdf import PdfReader, PdfWriter

    p = tmp_path / "spec.pdf"
    p.write_bytes(_minimal_pdf("Alpha motor specification"))
    r = ContentExtractor().extract(p)
    assert r.status == STATUS_OK and "Alpha" in r.text

    w = PdfWriter()
    for page in PdfReader(str(p)).pages:
        w.add_page(page)
    w.encrypt("secret")
    enc = tmp_path / "locked.pdf"
    with open(enc, "wb") as f:
        w.write(f)
    assert ContentExtractor().extract(enc).status == STATUS_ENCRYPTED


def test_office_formats(tmp_path):
    import docx
    import openpyxl
    from pptx import Presentation

    d = docx.Document()
    d.add_paragraph("Alpha 設計レビュー議事録")
    dp = tmp_path / "a.docx"
    d.save(dp)
    wb = openpyxl.Workbook()
    wb.active.append(["試験項目", "結果"])
    wb.active.append(["Alpha 耐久", "OK"])
    xp = tmp_path / "a.xlsx"
    wb.save(xp)
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Alpha 報告"
    pp = tmp_path / "a.pptx"
    prs.save(pp)
    ex = ContentExtractor()
    assert "議事録" in ex.extract(dp).text
    assert "Alpha 耐久" in ex.extract(xp).text
    assert "Alpha 報告" in ex.extract(pp).text


def test_password_protected_office_and_broken_zip(tmp_path):
    p = tmp_path / "locked.xlsx"
    p.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600)
    assert ContentExtractor().extract(p).status == STATUS_ENCRYPTED
    b = tmp_path / "broken.docx"
    b.write_bytes(b"not a zip")
    assert ContentExtractor().extract(b).status == STATUS_UNREADABLE
    z = tmp_path / "empty.docx"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("x.txt", "x")
    assert ContentExtractor().extract(z).status == STATUS_UNREADABLE


def test_image_metadata_and_unknown_binary(tmp_path):
    import struct
    png = tmp_path / "photo.png"
    png.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\rIHDR" + struct.pack(">II", 640, 480) + b"\x00" * 20)
    r = ContentExtractor().extract(png)
    assert r.status == STATUS_METADATA_ONLY and r.metadata["width"] == 640
    binf = tmp_path / "blob.bin"
    binf.write_bytes(b"\x00\x01\x02" * 100)
    assert ContentExtractor().extract(binf).status == STATUS_UNSUPPORTED


def test_access_denied_is_reported(tmp_path, monkeypatch):
    p = tmp_path / "denied.txt"
    p.write_text("x", encoding="utf-8")
    real_open = builtins.open

    def fake_open(file, *a, **k):
        if Path(str(file)) == p:
            raise PermissionError(13, "Permission denied")
        return real_open(file, *a, **k)

    monkeypatch.setattr(builtins, "open", fake_open)
    r = ContentExtractor().extract(p)
    assert r.status == STATUS_UNREADABLE and "拒否" in r.error
