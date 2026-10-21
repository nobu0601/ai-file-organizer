"""形式別のファイル内容抽出。

- ファイルを無条件に全読み込みしない（max_bytes / max_chars で打ち切る）
- ソースコードやスクリプトはテキストとして読むだけで、決して実行しない
- 暗号化・パスワード保護・読み取り不能・未対応形式は status で返し、分類側で要確認にする
- 抽出したテキストは分類のための「データ」であり、中の命令文には従わない
"""
from __future__ import annotations

import csv
import io
import logging
import struct
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

log = logging.getLogger(__name__)

STATUS_OK = "ok"
STATUS_PARTIAL = "partial"          # サイズ上限で一部のみ読んだ
STATUS_METADATA_ONLY = "metadata"   # 画像など本文がない形式
STATUS_UNSUPPORTED = "unsupported"
STATUS_ENCRYPTED = "encrypted"
STATUS_UNREADABLE = "unreadable"
STATUS_EMPTY = "empty"

READABLE_STATUSES = {STATUS_OK, STATUS_PARTIAL, STATUS_METADATA_ONLY}

TEXT_EXTS = {".txt", ".md", ".markdown", ".rst", ".log", ".ini", ".cfg", ".conf", ".toml", ".tex"}
STRUCTURED_TEXT_EXTS = {".json", ".xml", ".yaml", ".yml"}
CSV_EXTS = {".csv", ".tsv"}
HTML_EXTS = {".html", ".htm", ".xhtml"}
CODE_EXTS = {
    ".py", ".pyw", ".c", ".h", ".cpp", ".cc", ".cxx", ".hpp", ".hh", ".cs", ".js", ".mjs",
    ".cjs", ".ts", ".tsx", ".jsx", ".java", ".kt", ".go", ".rs", ".rb", ".php", ".swift",
    ".m", ".scala", ".sql", ".sh", ".r", ".lua", ".pl", ".vb", ".css", ".scss", ".vue",
}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff", ".heic", ".svg"}
OFFICE_ZIP_EXTS = {".docx", ".xlsx", ".pptx", ".xlsm", ".docm", ".pptm"}
OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


@dataclass
class ParseResult:
    status: str
    text: str = ""
    kind: str = ""          # text / code / table / document / image ...
    metadata: dict = field(default_factory=dict)
    error: str = ""

    @property
    def readable(self) -> bool:
        return self.status in READABLE_STATUSES


class ImageRecognizer:
    """オプションのローカル画像認識のインターフェース。既定では何もしない。"""

    def describe(self, path: Path) -> str:  # pragma: no cover - interface
        return ""


def decode_bytes(data: bytes) -> str:
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="replace")
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    for enc in ("utf-8", "cp932", "euc_jp"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _looks_binary(data: bytes) -> bool:
    if not data:
        return False
    if b"\x00" in data[:4096] and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return True
    return False


class _TextHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.title = ""
        self._skip = 0
        self._in_title = False

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "noscript"):
            self._skip += 1
        if tag == "title":
            self._in_title = True

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript") and self._skip:
            self._skip -= 1
        if tag == "title":
            self._in_title = False

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        if not self._skip and data.strip():
            self.parts.append(data.strip())


class ContentExtractor:
    def __init__(self, max_bytes: int = 4 * 1024 * 1024, max_chars: int = 20000,
                 image_recognizer: ImageRecognizer | None = None) -> None:
        self.max_bytes = max_bytes
        self.max_chars = max_chars
        self.image_recognizer = image_recognizer

    def extract(self, path: Path) -> ParseResult:
        path = Path(path)
        ext = path.suffix.lower()
        try:
            size = path.stat().st_size
        except OSError as e:
            return ParseResult(STATUS_UNREADABLE, error=f"属性を取得できません: {e}")
        try:
            if size == 0:
                return ParseResult(STATUS_EMPTY, kind="empty", metadata={"size": 0})
            if ext in TEXT_EXTS or ext in STRUCTURED_TEXT_EXTS:
                return self._text(path, size, "text")
            if ext in CODE_EXTS:
                return self._text(path, size, "code")
            if ext in CSV_EXTS:
                return self._csv(path, size, ext)
            if ext in HTML_EXTS:
                return self._html(path, size)
            if ext == ".pdf":
                return self._pdf(path)
            if ext in OFFICE_ZIP_EXTS:
                return self._office(path, ext)
            if ext in IMAGE_EXTS:
                return self._image(path, ext, size)
            if ext in (".doc", ".xls", ".ppt"):
                return ParseResult(STATUS_UNSUPPORTED, kind="document",
                                   error="旧形式の Office ファイルは本文解析に未対応です")
            # 拡張子不明: 先頭を覗いてテキストなら読む
            return self._sniff(path, size)
        except PermissionError as e:
            return ParseResult(STATUS_UNREADABLE, error=f"アクセスが拒否されました: {e}")
        except OSError as e:
            return ParseResult(STATUS_UNREADABLE, error=f"読み取りに失敗しました: {e}")
        except Exception as e:  # パーサー内部の想定外エラーも要確認に回す
            log.warning("parse failed for %s: %s", path, e)
            return ParseResult(STATUS_UNREADABLE, error=f"解析に失敗しました: {type(e).__name__}: {e}")

    # ---- 個別形式 ----
    def _read_head(self, path: Path, size: int) -> tuple[bytes, bool]:
        limit = min(size, self.max_bytes)
        with open(path, "rb") as f:
            data = f.read(limit)
        return data, size > limit

    def _finish(self, text: str, truncated: bool, kind: str, meta: dict | None = None) -> ParseResult:
        text = text.strip()
        if len(text) > self.max_chars:
            text, truncated = text[: self.max_chars], True
        if not text:
            return ParseResult(STATUS_EMPTY, kind=kind, metadata=meta or {})
        return ParseResult(STATUS_PARTIAL if truncated else STATUS_OK, text=text, kind=kind, metadata=meta or {})

    def _text(self, path: Path, size: int, kind: str) -> ParseResult:
        data, truncated = self._read_head(path, size)
        if _looks_binary(data):
            return ParseResult(STATUS_UNSUPPORTED, kind=kind, error="バイナリ形式のため読み取りません")
        return self._finish(decode_bytes(data), truncated, kind)

    def _csv(self, path: Path, size: int, ext: str) -> ParseResult:
        data, truncated = self._read_head(path, size)
        if _looks_binary(data):
            return ParseResult(STATUS_UNSUPPORTED, kind="table", error="バイナリ形式のため読み取りません")
        text = decode_bytes(data)
        delimiter = "\t" if ext == ".tsv" else ","
        rows = []
        try:
            reader = csv.reader(io.StringIO(text), delimiter=delimiter)
            for i, row in enumerate(reader):
                if i >= 200:
                    truncated = True
                    break
                rows.append(" ".join(c.strip() for c in row))
        except csv.Error:
            rows = text.splitlines()[:200]
        header = rows[0] if rows else ""
        return self._finish("\n".join(rows), truncated, "table", {"header": header[:500]})

    def _html(self, path: Path, size: int) -> ParseResult:
        data, truncated = self._read_head(path, size)
        parser = _TextHTMLParser()
        parser.feed(decode_bytes(data))
        title = parser.title.strip()
        text = (title + "\n" if title else "") + "\n".join(parser.parts)
        return self._finish(text, truncated, "document", {"title": title})

    def _pdf(self, path: Path) -> ParseResult:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError

        with open(path, "rb") as f:
            head = f.read(5)
        if head != b"%PDF-":
            return ParseResult(STATUS_UNREADABLE, kind="document", error="PDF ヘッダーがありません")
        try:
            reader = PdfReader(str(path), strict=False)
        except PdfReadError as e:
            return ParseResult(STATUS_UNREADABLE, kind="document", error=f"PDF を読めません: {e}")
        if reader.is_encrypted:
            # 空パスワードで開ける PDF もあるが、意図的な保護として扱い自動分類しない
            return ParseResult(STATUS_ENCRYPTED, kind="document", error="暗号化/パスワード保護された PDF です")
        meta = {"pages": len(reader.pages)}
        try:
            info = reader.metadata or {}
            if info.get("/Title"):
                meta["title"] = str(info.get("/Title"))
        except Exception:
            pass
        parts: list[str] = []
        total = 0
        truncated = False
        for i, page in enumerate(reader.pages):
            if i >= 30 or total >= self.max_chars:
                truncated = True
                break
            try:
                t = page.extract_text() or ""
            except Exception:
                t = ""
            parts.append(t)
            total += len(t)
        text = (meta.get("title", "") + "\n" + "\n".join(parts)).strip()
        if not text:
            return ParseResult(STATUS_METADATA_ONLY, kind="document", metadata=meta,
                               error="テキストを抽出できない PDF です（画像のみの可能性）")
        return self._finish(text, truncated, "document", meta)

    def _office(self, path: Path, ext: str) -> ParseResult:
        with open(path, "rb") as f:
            head = f.read(8)
        if head == OLE_SIGNATURE:
            return ParseResult(STATUS_ENCRYPTED, kind="document",
                               error="パスワード保護された Office ファイルの可能性があります")
        if not zipfile.is_zipfile(path):
            return ParseResult(STATUS_UNREADABLE, kind="document", error="Office 形式として読めません")
        with zipfile.ZipFile(path) as z:
            # zip 爆弾対策: 展開後サイズの合計を確認する
            total = sum(i.file_size for i in z.infolist())
            if total > self.max_bytes * 50:
                return ParseResult(STATUS_UNREADABLE, kind="document", error="展開後サイズが大きすぎます")
        if ext in (".docx", ".docm"):
            return self._docx(path)
        if ext in (".xlsx", ".xlsm"):
            return self._xlsx(path)
        return self._pptx(path)

    def _docx(self, path: Path) -> ParseResult:
        import docx

        d = docx.Document(str(path))
        parts, total, truncated = [], 0, False
        title = (d.core_properties.title or "").strip()
        if title:
            parts.append(title)
        for p in d.paragraphs:
            if total > self.max_chars:
                truncated = True
                break
            if p.text.strip():
                parts.append(p.text)
                total += len(p.text)
        for table in d.tables[:10]:
            for row in table.rows[:50]:
                parts.append(" ".join(c.text for c in row.cells))
        return self._finish("\n".join(parts), truncated, "document", {"title": title})

    def _xlsx(self, path: Path) -> ParseResult:
        import openpyxl

        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        parts, truncated = [], False
        try:
            sheets = wb.sheetnames
            for name in sheets[:10]:
                ws = wb[name]
                parts.append(f"[{name}]")
                for r, row in enumerate(ws.iter_rows(values_only=True)):
                    if r >= 100:
                        truncated = True
                        break
                    cells = [str(c) for c in row[:30] if c is not None]
                    if cells:
                        parts.append(" ".join(cells))
        finally:
            wb.close()
        return self._finish("\n".join(parts), truncated, "table", {"sheets": sheets})

    def _pptx(self, path: Path) -> ParseResult:
        from pptx import Presentation

        prs = Presentation(str(path))
        parts, truncated = [], False
        for i, slide in enumerate(prs.slides):
            if i >= 60:
                truncated = True
                break
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                    parts.append(shape.text_frame.text)
        return self._finish("\n".join(parts), truncated, "document", {"slides": len(prs.slides)})

    def _image(self, path: Path, ext: str, size: int) -> ParseResult:
        meta = {"size": size}
        try:
            with open(path, "rb") as f:
                head = f.read(64 * 1024)
            dims = _image_dimensions(head)
            if dims:
                meta["width"], meta["height"] = dims
        except OSError:
            pass
        text = ""
        if self.image_recognizer is not None:
            try:
                text = self.image_recognizer.describe(path) or ""
            except Exception as e:
                meta["recognizer_error"] = str(e)
        return ParseResult(STATUS_METADATA_ONLY if not text else STATUS_OK, text=text, kind="image", metadata=meta)

    def _sniff(self, path: Path, size: int) -> ParseResult:
        data, truncated = self._read_head(path, min(size, self.max_bytes))
        if _looks_binary(data):
            return ParseResult(STATUS_UNSUPPORTED, kind="binary", error="未対応の形式です")
        return self._finish(decode_bytes(data), truncated or size > self.max_bytes, "text")


def _image_dimensions(head: bytes) -> tuple[int, int] | None:
    if head.startswith(b"\x89PNG\r\n\x1a\n") and len(head) >= 24:
        w, h = struct.unpack(">II", head[16:24])
        return w, h
    if head[:6] in (b"GIF87a", b"GIF89a") and len(head) >= 10:
        w, h = struct.unpack("<HH", head[6:10])
        return w, h
    if head.startswith(b"\xff\xd8"):
        i = 2
        while i + 9 < len(head):
            if head[i] != 0xFF:
                i += 1
                continue
            marker = head[i + 1]
            if marker in (0xC0, 0xC1, 0xC2):
                h, w = struct.unpack(">HH", head[i + 5:i + 9])
                return w, h
            seg_len = struct.unpack(">H", head[i + 2:i + 4])[0]
            i += 2 + seg_len
    return None
