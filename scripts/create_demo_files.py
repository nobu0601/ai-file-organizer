"""動作確認用のデモファイルを指定フォルダに作成する（既存ファイルは上書きしない）。

    python scripts/create_demo_files.py C:\\Users\\you\\Desktop\\demo-downloads
"""
from __future__ import annotations

import sys
from pathlib import Path

TEXT_FILES = {
    "alpha_設計書_v2.md": "# Alphaプロジェクト 設計書\nモーター制御基板の仕様と試験計画を記載する。",
    "alpha_試験結果.csv": "項目,結果\nAlpha モーター 耐久試験,合格\nAlpha 温度試験,合格\n",
    "alpha_解析メモ.txt": "Alpha 解析: モーター電流の解析結果。Alpha基板の温度特性をまとめた。",
    "請求書_2026_04_山田商事.txt": "請求書 山田商事 御中 金額 120,000円 支払期限 2026年5月末 振込先 ○○銀行",
    "請求書_2026_05_山田商事.txt": "請求書 山田商事 御中 金額 80,000円 支払期限 2026年6月末 振込先 ○○銀行",
    "travel_kyoto_plan.md": "Kyoto trip itinerary: temple visit, hotel booking, train schedule",
    "kyoto_hotel_booking.html": "<html><title>Kyoto Hotel</title><body>Hotel booking confirmation Kyoto itinerary</body></html>",
    "sensor_logger.py": "# 温度センサーのロガー（このアプリは実行しません）\nimport time\nprint('sensor log')\n",
    "secret_memo.txt": "サーバーの password はここに書かない",
}


def main(target: str) -> None:
    root = Path(target)
    root.mkdir(parents=True, exist_ok=True)
    made = 0
    for name, body in TEXT_FILES.items():
        p = root / name
        if not p.exists():
            p.write_text(body, encoding="utf-8")
            made += 1
    try:
        import openpyxl
        p = root / "alpha_試験データ.xlsx"
        if not p.exists():
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["Alpha 試験項目", "結果", "温度"])
            ws.append(["モーター 耐久", "OK", 45])
            wb.save(p)
            made += 1
        import docx
        p = root / "Alpha_設計レビュー議事録.docx"
        if not p.exists():
            d = docx.Document()
            d.add_paragraph("Alpha 設計レビュー議事録: モーター基板の試験計画を確認した。")
            d.save(p)
            made += 1
    except ImportError:
        pass
    print(f"{made} 件のデモファイルを作成しました: {root}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1])
