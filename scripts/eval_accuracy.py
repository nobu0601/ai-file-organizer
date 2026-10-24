"""正解ラベル付きデータで分類精度を測る（開発・調整用）。

    python scripts/eval_accuracy.py                       # 内容ベース分類のみ
    python scripts/eval_accuracy.py --models qwen2.5:3b   # Ollama のモデルも比較
    python scripts/eval_accuracy.py --models qwen2.5:1.5b,qwen2.5:3b,qwen2.5:7b --out report.md

すべて一時ディレクトリ内で行い、実際のファイルには触れません。

測るもの:
  1. 分類体系の考案: seed ファイル群から作ったカテゴリが正解テーマとどれだけ一致するか
     （純度・ペア F1・カテゴリが作られたテーマ数・カテゴリに入らなかった件数）
  2. 新規ファイルの分類: 承認後に holdout（正解テーマあり）と distractor（どのテーマでもない）を処理
     - 正解率（holdout のうち正しいカテゴリに分類された割合。名前に手掛かりのないファイルは別集計）
     - 自動移動の精度（自動移動されたもののうち正しかった割合）と件数
     - distractor の誤移動件数（0 であるべき）
  3. 処理時間
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
import tempfile
import time
from collections import Counter, defaultdict
from itertools import combinations
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from classifier.heuristic import HeuristicClassifier  # noqa: E402
from classifier.ollama import OllamaClassifier, OllamaClient  # noqa: E402
from classifier.semantic import SemanticClassifier  # noqa: E402
from organizer.engine import Engine  # noqa: E402
from settings.config import Settings, WatchFolder  # noqa: E402
from settings.paths import AppPaths  # noqa: E402

GENERIC_NAME = re.compile(r"^(scan|document|無題|memo|新しい|export|img|untitled|data|file|download|コピー)",
                          re.IGNORECASE)


def write_file(path: Path, content: str) -> None:
    ext = path.suffix.lower()
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.Workbook()
        for row in csv.reader(io.StringIO(content)):
            wb.active.append(row)
        wb.save(path)
    elif ext == ".docx":
        import docx
        d = docx.Document()
        for line in content.splitlines():
            if line.strip():
                d.add_paragraph(line)
        d.save(path)
    else:
        path.write_text(content, encoding="utf-8")


def pair_f1(assign: dict[str, str | None], gold: dict[str, str]) -> tuple[float, float, float]:
    """同じカテゴリに入ったペアと、正解で同じテーマのペアの一致度。"""
    names = sorted(gold)
    tp = fp = fn = 0
    for a, b in combinations(names, 2):
        same_pred = assign.get(a) is not None and assign.get(a) == assign.get(b)
        same_gold = gold[a] == gold[b]
        tp += same_pred and same_gold
        fp += same_pred and not same_gold
        fn += same_gold and not same_pred
    p = tp / (tp + fp) if tp + fp else 0.0
    r = tp / (tp + fn) if tp + fn else 0.0
    return p, r, (2 * p * r / (p + r) if p + r else 0.0)


def run(dataset: dict, label: str, classifier, threshold: float) -> dict:
    files = dataset["files"]
    with tempfile.TemporaryDirectory(prefix="aifo-eval-") as tmp:
        tmp = Path(tmp)
        watch, dest, data = tmp / "Downloads", tmp / "Organized", tmp / "appdata"
        watch.mkdir()
        s = Settings(organize_root=str(dest), watch_folders=[WatchFolder(str(watch), False)],
                     preview_mode=False, auto_move=True, confidence_threshold=threshold)
        engine = Engine(AppPaths(data), s, classifier=classifier)
        try:
            seed = [f for f in files if f["split"] == "seed"]
            for f in seed:
                write_file(watch / f["name"], f["content"])
            t0 = time.time()
            summary = engine.propose_taxonomy()
            t_propose = time.time() - t0
            props = engine.db.list_proposals("pending")
            gold_seed = {f["name"]: f["label"] for f in seed}
            assign: dict[str, str | None] = {n: None for n in gold_seed}
            cat_members: dict[str, list[str]] = defaultdict(list)
            cat_names: dict[str, str] = {}
            for p in props:
                cid = p.operations[0]["category"]["category_id"]
                cat_names[cid] = p.operations[0]["category"]["category_name"]
                for path in p.related_files:
                    n = Path(path).name
                    assign[n] = cid
                    cat_members[cid].append(n)
            # 子カテゴリの所属は親にも数えない（各ファイルは最も具体的なカテゴリに属する）
            cat_label = {cid: Counter(gold_seed[m] for m in ms).most_common(1)[0][0]
                         for cid, ms in cat_members.items() if ms}
            assigned = [n for n, c in assign.items() if c]
            purity = (sum(Counter(gold_seed[m] for m in ms).most_common(1)[0][1] for ms in cat_members.values())
                      / len(assigned)) if assigned else 0.0
            pp, pr, pf = pair_f1(assign, gold_seed)
            themes_covered = len(set(cat_label.values()))

            ok, msg = engine.approve_proposals([p.id for p in props])
            # 承認時の再分類で seed が実行待ちになる → 移動して整理済みの状態にする
            engine.execute_ready()

            rows = []
            t1 = time.time()
            for f in files:
                if f["split"] == "seed":
                    continue
                p = watch / f["name"]
                write_file(p, f["content"])
                out = engine.process_file(p, "event")
                cat = engine.store.taxonomy.get(out.category_id)
                pred = None
                if cat:
                    # 下位カテゴリが正解テーマを持たない場合は親をたどる
                    c = cat
                    while c and c.category_id not in cat_label and c.parent_category_id:
                        c = engine.store.taxonomy.get(c.parent_category_id)
                    pred = cat_label.get(c.category_id) if c else None
                rows.append({
                    "name": f["name"], "gold": f["label"], "split": f["split"], "pred": pred,
                    "category": cat.relative_path if cat else None, "status": out.status,
                    "confidence": out.confidence, "generic": bool(GENERIC_NAME.match(f["name"])),
                    "message": out.message,
                })
            t_classify = time.time() - t1
        finally:
            engine.shutdown()

    hold = [r for r in rows if r["split"] == "holdout"]
    dist = [r for r in rows if r["split"] == "distractor"]
    gen = [r for r in hold if r["generic"]]
    desc = [r for r in hold if not r["generic"]]
    acc = lambda rs: (sum(r["pred"] == r["gold"] for r in rs) / len(rs)) if rs else 0.0  # noqa: E731
    moved = [r for r in rows if r["status"] == "moved"]
    moved_ok = [r for r in moved if r["pred"] == r["gold"]]
    return {
        "label": label, "approve": msg if ok else f"承認失敗: {msg}",
        "propose_seconds": round(t_propose, 1), "classify_seconds": round(t_classify, 1),
        "n_files_classified": len(rows), "categories": len(props), "category_names": sorted(cat_names.values()),
        "unassigned": summary["unassigned"], "purity": purity, "pair_precision": pp, "pair_recall": pr,
        "pair_f1": pf, "themes_covered": themes_covered, "n_themes": len(dataset["themes"]),
        "holdout_acc": acc(hold), "holdout_acc_generic": acc(gen), "holdout_acc_named": acc(desc),
        "n_generic": len(gen), "n_named": len(desc),
        "auto_moved": len(moved), "auto_moved_correct": len(moved_ok),
        "auto_move_precision": (len(moved_ok) / len(moved)) if moved else None,
        "auto_move_coverage": len([r for r in hold if r["status"] == "moved" and r["pred"] == r["gold"]]) / len(hold)
        if hold else 0.0,
        "distractor_moved": sum(r["status"] == "moved" for r in dist), "n_distractors": len(dist),
        "rows": rows,
    }


def pct(x) -> str:
    return "—" if x is None else f"{x * 100:.0f}%"


def report(results: list[dict]) -> str:
    lines = ["# 分類精度の評価結果", "",
             "## 1. 分類体系の考案（seed ファイルから）", "",
             "| 方式 | カテゴリ数 | 作られたテーマ | 純度 | ペアF1 | 未分類 | 所要時間 |",
             "|---|---|---|---|---|---|---|"]
    for r in results:
        lines.append(f"| {r['label']} | {r['categories']} | {r['themes_covered']}/{r['n_themes']} | {pct(r['purity'])} | "
                     f"{pct(r['pair_f1'])} | {r['unassigned']} | {r['propose_seconds']}秒 |")
    lines += ["", "## 2. 新規ファイルの分類（承認後）", "",
              "| 方式 | 正解率 | 名前あり | 名前に手掛かりなし | 自動移動（正しい/件数） | 自動移動の精度 | 無関係ファイルの誤移動 | 所要時間 |",
              "|---|---|---|---|---|---|---|---|"]
    for r in results:
        lines.append(
            f"| {r['label']} | {pct(r['holdout_acc'])} | {pct(r['holdout_acc_named'])} ({r['n_named']}件) | "
            f"{pct(r['holdout_acc_generic'])} ({r['n_generic']}件) | {r['auto_moved_correct']}/{r['auto_moved']} | "
            f"{pct(r['auto_move_precision'])} | {r['distractor_moved']}/{r['n_distractors']} | {r['classify_seconds']}秒 |")
    for r in results:
        lines += ["", f"### {r['label']} の詳細", "", f"カテゴリ: {', '.join(r['category_names'])}", "",
                  "| ファイル | 正解 | 予測 | 状態 | 信頼度 |", "|---|---|---|---|---|"]
        for row in r["rows"]:
            mark = "✓" if row["pred"] == row["gold"] else ("—" if row["gold"] == "none" and not row["pred"] else "✗")
            conf = f"{row['confidence']:.2f}" if row["confidence"] is not None else ""
            lines.append(f"| {row['name']} | {row['gold']} | {row['pred'] or '—'} {mark} | {row['status']} | {conf} |")
    return "\n".join(lines) + "\n"


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dataset", default=str(ROOT / "scripts" / "eval" / "dataset.json"))
    ap.add_argument("--models", default="", help="カンマ区切りの Ollama モデル名")
    ap.add_argument("--url", default="http://127.0.0.1:11434")
    ap.add_argument("--timeout", type=float, default=300)
    ap.add_argument("--threshold", type=float, default=0.85)
    ap.add_argument("--no-heuristic", action="store_true")
    ap.add_argument("--embed", default="", help="埋め込みモデル名（例: bge-m3）。--models と併用すると LLM が命名する")
    ap.add_argument("--out", help="Markdown レポートの出力先")
    ap.add_argument("--json", help="詳細結果（JSON）の出力先")
    args = ap.parse_args()
    dataset = json.loads(Path(args.dataset).read_text(encoding="utf-8"))
    results = []
    if not args.no_heuristic:
        print("内容ベース分類を評価中…", flush=True)
        results.append(run(dataset, "内容ベース（LLM なし）", HeuristicClassifier(), args.threshold))
    models = [x.strip() for x in args.models.split(",") if x.strip()]
    if args.embed:
        # 埋め込みのみ（命名は内容ベース）と、埋め込み＋各 LLM（命名を LLM）を評価する
        for m in [None] + models:
            client = OllamaClient(args.url, m or "auto", args.timeout)
            if not client.has_model(args.embed):
                print(f"{args.embed}: スキップ（インストールされていません）", flush=True)
                break
            namer = OllamaClassifier(client) if m else None
            label = f"埋め込み {args.embed}" + (f" + 命名 {m}" if m else "")
            print(f"{label} を評価中…", flush=True)
            clf = SemanticClassifier(client, args.embed, namer=namer)
            res = run(dataset, label, clf, args.threshold)
            if clf.last_error:
                res["label"] += "（途中でエラーあり）"
            results.append(res)
        models = []
    for m in models:
        client = OllamaClient(args.url, m, args.timeout)
        ok, msg = client.available()
        if not ok:
            print(f"{m}: スキップ（{msg}）", flush=True)
            continue
        print(f"Ollama {m} を評価中…", flush=True)
        clf = OllamaClassifier(client)
        res = run(dataset, f"Ollama {m}", clf, args.threshold)
        if clf.last_error:
            res["label"] += "（途中で LLM エラーあり）"
        results.append(res)
    text = report(results)
    print(text)
    if args.out:
        Path(args.out).write_text(text, encoding="utf-8")
    if args.json:
        Path(args.json).write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0


if __name__ == "__main__":
    sys.exit(main())
